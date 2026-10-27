"""
El informe Word, con dos corridas sintéticas (sin modelos ni videos): se prueba
que se arma con datos parciales y completos, y que la comparación pareada sale
de las mismas filas.
"""

import csv
import json
from pathlib import Path

import pytest

pytest.importorskip("docx")
pytest.importorskip("pandas")
pytest.importorskip("matplotlib")

from docx import Document  # noqa: E402

from benchmark import results  # noqa: E402
from scripts import build_report  # noqa: E402

MODELS = ["yolo26n", "mediapipe_lite"]
VIDEOS = ["V1", "V2"]


def _write_run(directory: Path, environment: str, ms_per_frame: dict, sessions: int | None = None) -> Path:
    """Una corrida con 3 repeticiones medidas por modelo-video; `sessions` corta las primeras N."""
    directory.mkdir(parents=True)
    rows = []
    combos = [(m, v) for v in VIDEOS for m in MODELS][:sessions]
    for model, video in combos:
        for repetition in (1, 2, 3):
            processing = ms_per_frame[model] * 300 / 1000 * (1 + 0.01 * repetition)
            rows.append({
                "run_id": directory.name, "environment": environment, "model": model, "video_id": video,
                "phase": "measured", "repetition": repetition, "trial": 1, "concurrency": 1,
                "processing_s": processing, "job_s": processing + 0.05, "model_init_s": 0.05,
                "ms_per_frame": processing * 1000 / 300, "realtime_factor": processing / 10, "fps_processed": 300 / processing,
                "frames_processed": 300, "video_frames_reported": 300, "video_fps": 30, "video_duration_s": 10,
                "video_size_mb": 20, "video_stored_resolution": "1920x1080", "video_decoded_resolution": "1080x1920",
                "video_rotation_meta_deg": 90, "cpu_cores_used": 1.5, "cpu_s": 1.5 * processing,
                "ram_mb_peak": 400, "ram_start_mb": 100 + repetition,
                "detection_rate": 1.0, "decode_ms_mean": 7.0, "infer_ms_p90": ms_per_frame[model],
                "cgroup_throttled_pct": 10, "cpu_model": "CPU de prueba", "cpu_limit_cores": 2, "memory_limit_mb": 8192,
                "threads_requested": 2, "versions_match_service": True,
            })
    with (directory / "rows.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=results.ROW_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    info = {"run_id": directory.name, "models": MODELS, "videos": [{"id": v} for v in VIDEOS],
            "settings": {"repetitions": 3, "trials": 1}, "cpu_model": "CPU de prueba"}
    (directory / "run_info.json").write_text(json.dumps(info), encoding="utf-8")
    return directory


def _text(path: Path) -> str:
    document = Document(str(path))
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


def test_report_with_only_a_partial_pc_run_says_what_is_missing(tmp_path):
    pc = _write_run(tmp_path / "pc", "docker", {"yolo26n": 80, "mediapipe_lite": 25}, sessions=3)
    out = build_report.build(build_report.Run("PC", pc), build_report.Run("Render", tmp_path / "no_existe"),
                             tmp_path / "informe.docx", None)

    text = _text(out)
    assert "en curso: 3 de 4 sesiones" in text
    assert "Render: sin datos todavía" in text
    assert "Pendiente" in text  # la sección de diferencias espera a Render


def test_report_with_both_runs_compares_them_pairwise(tmp_path):
    pc = _write_run(tmp_path / "pc", "docker", {"yolo26n": 80, "mediapipe_lite": 25})
    render = _write_run(tmp_path / "render", "render", {"yolo26n": 104, "mediapipe_lite": 30})
    out = build_report.build(build_report.Run("PC", pc), build_report.Run("Render", render), tmp_path / "informe.docx", None)

    text = _text(out)
    assert "Estado del informe: final." in text
    assert "×1,30" in text  # YOLO: 104 / 80 en cada video
    assert "+30 %" in text
    assert (tmp_path / "informe_figuras" / "02_render_vs_pc.png").is_file()


def test_concurrency_2_run_gets_its_own_section_and_notes_are_kept(tmp_path):
    pc = _write_run(tmp_path / "pc", "docker", {"yolo26n": 80, "mediapipe_lite": 25})
    render = _write_run(tmp_path / "render", "render", {"yolo26n": 104, "mediapipe_lite": 30})
    render_c2 = _write_run(tmp_path / "render_c2", "render", {"yolo26n": 150, "mediapipe_lite": 40})
    out = build_report.build(build_report.Run("PC", pc), build_report.Run("Render", render), tmp_path / "informe.docx",
                             None, build_report.Run(build_report.RENDER_C2, render_c2), ["la repetición 3 se perturbó"])

    text = _text(out)
    assert "5. Render con dos videos a la vez" in text
    assert "Videos/hora, 2" in text
    assert "Nota. la repetición 3 se perturbó" in text
    assert (tmp_path / "informe_figuras" / "05a_videos_por_hora.png").is_file()


def test_worker_ram_adds_the_processes_that_ran_together():
    run = build_report.Run("x", None)
    run.rows = [{"phase": "measured", "model": "m", "video_id": "V", "trial": "1", "repetition": "1", "ram_mb_peak": "400"},
                {"phase": "measured", "model": "m", "video_id": "V", "trial": "1", "repetition": "1", "ram_mb_peak": "450"},
                {"phase": "measured", "model": "m", "video_id": "V", "trial": "1", "repetition": "2", "ram_mb_peak": "300"}]

    assert build_report.worker_ram_by_model(run) == {"m": 850.0}


def test_paired_ratio_is_render_over_pc():
    import pandas as pd

    summary = pd.DataFrame([
        {"environment": "PC", "model": "m", "video_id": "V", "ms_per_frame_median": 50.0, "cpu_cores_used_median": 1.0, "ram_mb_peak_max": 300.0},
        {"environment": "Render", "model": "m", "video_id": "V", "ms_per_frame_median": 65.0, "cpu_cores_used_median": 1.1, "ram_mb_peak_max": 310.0},
    ])
    pairs = build_report.paired(summary)

    assert pairs.loc[0, "ratio_ms"] == pytest.approx(1.3)
    assert pairs.loc[0, "ram_diff_mb"] == pytest.approx(10)


def test_numbers_use_decimal_comma():
    assert build_report.num(1234.5, 1) == "1.234,5"
    assert build_report.num(None) == "n/d"
    assert build_report.pct_change(0.9) == "−10 %"
