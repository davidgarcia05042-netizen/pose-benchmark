"""
Informe comparativo en Word: la misma suite medida en Render y en el PC.

Análisis LOCAL (no va en la imagen). Lee las carpetas `results/<run_id>/` de
las dos corridas (rows.csv + run_info.json, como las deja run_benchmark o
collect_from_log), resume con benchmark/stats.py, dibuja las gráficas y arma
un .docx con tablas, gráficas, diferencias y limitaciones.

El PC es la línea base (Docker con 2 CPU / 8 GB, la réplica del presupuesto) y
Render es el entorno que decide. Cualquiera de las dos corridas puede estar
incompleta o faltar: el informe dice qué tiene y qué falta, así que se puede
regenerar mientras las corridas avanzan.

Uso:
    python -m scripts.build_report \\
        --pc results/20261001_toma_frontal_docker_c1 \\
        --render results/20261001T063552Z \\
        --render-c2 results/20261001T070703Z \\
        --out results/informes/informe_render_vs_pc.docx \\
        --note "Texto sobre algo que pasó durante una corrida (se puede repetir)"

--render-c2 es la corrida de Render con --concurrency 2 (como el worker): se
compara solo contra Render con 1 a la vez, en su propia sección.
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from statistics import median

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from benchmark import config, stats
from scripts.analyze_results import load_rows
from scripts.charts import (MODEL_COLORS, geometric_mean, model_label, num, ordered_models, ratio_dots,
                            source_bars, video_bars)

PC, RENDER = "PC", "Render"
# El entorno es la identidad en las gráficas de comparación. No se usa ningún
# color de modelo (que ya tienen dueño en las gráficas por video): el PC, la
# línea base, va en gris; Render en violeta (ranura 7 de la paleta categórica).
SOURCE_COLORS = {PC: "#8a8984", RENDER: "#4a3aa7"}
# La corrida de Render con --concurrency 2 (como el worker). Solo se compara
# contra Render con 1 a la vez, en su propia sección: mezclarla con el PC
# confundiría el efecto del procesador con el de la concurrencia.
RENDER_C2 = "Render ×2"
CONCURRENCY_COLORS = {RENDER: "#4a3aa7", RENDER_C2: "#e87ba4"}
VERDICT_RANK = {stats.ACCEPTABLE: 0, stats.CONDITIONAL: 1, stats.NOT_ACCEPTABLE: 2}
VERDICT_LABELS = {stats.ACCEPTABLE: "aceptable", stats.CONDITIONAL: "condicionado", stats.NOT_ACCEPTABLE: "no aceptable"}
HEADER_FILL = "E4E3DF"
NOTE_FILL = "F3F2EE"


# --- Formato ---------------------------------------------------------------


def pct_change(ratio: float | None) -> str:
    if ratio is None:
        return "n/d"
    change = (ratio - 1) * 100
    return f"{'+' if change >= 0 else '−'}{num(abs(change), 0)} %"


# --- Datos -----------------------------------------------------------------


class Run:
    """Una corrida: sus filas, su run_info y cuánto le falta."""

    def __init__(self, label: str, directory: Path | None):
        self.label = label
        self.directory = directory
        self.rows: list[dict] = []
        self.info: dict = {}
        if directory and (directory / "rows.csv").is_file():
            self.rows = load_rows(directory / "rows.csv")
            for row in self.rows:
                # El entorno de la fila (docker_2cpu_8gb, render) se reemplaza por
                # la etiqueta de la fuente: así stats.summarize agrupa por fuente.
                row["environment"] = label
        if directory and (directory / "run_info.json").is_file():
            self.info = json.loads((directory / "run_info.json").read_text(encoding="utf-8"))

    @property
    def available(self) -> bool:
        return bool(self.rows)

    def expected_sessions(self) -> int | None:
        if not self.info:
            return None
        settings = self.info.get("settings", {})
        return len(self.info.get("videos", [])) * len(self.info.get("models", [])) * int(settings.get("trials") or 1)

    def completed_sessions(self) -> int:
        repetitions = int((self.info.get("settings") or {}).get("repetitions") or stats.MIN_REPETITIONS)
        counts: dict[tuple, int] = {}
        for row in self.rows:
            if row.get("phase") == "measured":
                key = (row.get("model"), row.get("video_id"), row.get("trial"))
                counts[key] = counts.get(key, 0) + 1
        return sum(1 for count in counts.values() if count >= repetitions)

    def status(self) -> str:
        if not self.available:
            return "sin datos todavía"
        expected = self.expected_sessions()
        done = self.completed_sessions()
        if expected and done < expected:
            return f"en curso: {done} de {expected} sesiones modelo-video completas"
        return f"completa: {done} sesiones modelo-video" + (f" de {expected}" if expected else "")

    def complete(self) -> bool:
        expected = self.expected_sessions()
        return self.available and expected is not None and self.completed_sessions() >= expected

    def first(self, column: str):
        for row in self.rows:
            if row.get(column) not in (None, ""):
                return row[column]
        return None


def summarize(runs: list[Run]) -> pd.DataFrame:
    rows = [row for run in runs for row in run.rows]
    summary = pd.DataFrame(stats.summarize(rows))
    if summary.empty:
        return summary
    # Solo sesiones completas: una a medias tiene menos repeticiones y ensucia la comparación.
    return summary[summary["n"] >= stats.MIN_REPETITIONS].reset_index(drop=True)


def by_model(summary: pd.DataFrame) -> pd.DataFrame:
    """Una fila por (fuente, modelo): agregados sobre los videos."""
    records = []
    for (source, model), group in summary.groupby(["environment", "model"]):
        verdicts = [v for v in group["verdict_time"] if v in VERDICT_RANK]
        ram_verdicts = [v for v in group["verdict_ram"] if v in VERDICT_RANK]
        records.append({
            "source": source,
            "model": model,
            "videos": len(group),
            "ms_per_frame": median(group["ms_per_frame_median"]),
            "ms_per_frame_min": group["ms_per_frame_median"].min(),
            "ms_per_frame_max": group["ms_per_frame_median"].max(),
            "realtime_factor": median(group["realtime_factor_median"]),
            "job_s_max": group["job_s_max"].max(),
            "model_init_s": median(group["model_init_s_median"]),
            "cores": median(group["cpu_cores_used_median"]),
            "throttled_pct": median(group["cgroup_throttled_pct_median"].dropna()) if group["cgroup_throttled_pct_median"].notna().any() else None,
            "ram_peak": group["ram_mb_peak_max"].max(),
            "ram_drift": group["ram_drift_mb"].max(),
            "cv_max": group["processing_s_cv_pct"].max(),
            "detection_min": group["detection_rate_min"].min(),
            # El video más largo que cabe en el apagado de 300 s, en el peor video.
            "max_video_s": group["max_video_s_within_shutdown"].min(),
            "throughput": median(group["throughput_videos_per_hour"]),
            "verdict_time": max(verdicts, key=VERDICT_RANK.get) if verdicts else None,
            "verdict_ram": max(ram_verdicts, key=VERDICT_RANK.get) if ram_verdicts else None,
            "valid_all": bool(group["measurement_valid"].all()),
        })
    return pd.DataFrame(records)


def paired(summary: pd.DataFrame) -> pd.DataFrame:
    """Render frente al PC en cada (modelo, video) medido en los dos."""
    if summary.empty:
        return pd.DataFrame()
    keys = ["model", "video_id"]
    columns = keys + ["ms_per_frame_median", "cpu_cores_used_median", "ram_mb_peak_max", "decode_ms_mean",
                      "model_init_s_median", "processing_s_cv_pct", "cgroup_throttled_pct_median"]
    pc = summary[summary["environment"] == PC][[c for c in columns if c in summary.columns]]
    render = summary[summary["environment"] == RENDER][[c for c in columns if c in summary.columns]]
    merged = pc.merge(render, on=keys, suffixes=("_pc", "_render"))
    if merged.empty:
        return merged
    merged["ratio_ms"] = merged["ms_per_frame_median_render"] / merged["ms_per_frame_median_pc"]
    merged["ram_diff_mb"] = merged["ram_mb_peak_max_render"] - merged["ram_mb_peak_max_pc"]
    merged["cores_diff"] = merged["cpu_cores_used_median_render"] - merged["cpu_cores_used_median_pc"]
    return merged


def decode_by_source(runs: list[Run]) -> dict[str, float]:
    """Mediana de decode_ms_mean (decodificar + rotar) por fuente: no depende del modelo."""
    result = {}
    for run in runs:
        values = [float(r["decode_ms_mean"]) for r in run.rows if r.get("phase") == "measured" and r.get("decode_ms_mean")]
        if values:
            result[run.label] = median(values)
    return result


# --- Gráficas --------------------------------------------------------------

# --- Word ------------------------------------------------------------------


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


class Report:
    def __init__(self):
        self.doc = Document()
        self.figure_number = 0
        self.table_number = 0
        section = self.doc.sections[0]
        # Carta, como pide la norma de presentación de trabajos en Colombia.
        section.page_width, section.page_height = Cm(21.59), Cm(27.94)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.5))
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)

    @property
    def text_width_cm(self) -> float:
        section = self.doc.sections[-1]
        return (section.page_width - section.left_margin - section.right_margin) / 360000

    def heading(self, text: str, level: int = 1) -> None:
        self.doc.add_heading(text, level=level)

    def para(self, text: str = "", bold_prefix: str | None = None, italic: bool = False, size: float | None = None):
        paragraph = self.doc.add_paragraph()
        if bold_prefix:
            paragraph.add_run(bold_prefix).bold = True
        run = paragraph.add_run(text)
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        return paragraph

    def bullets(self, items: list[str]) -> None:
        """Viñetas; lo que va entre **dobles asteriscos** sale en negrita, en cualquier parte del texto."""
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Bullet")
            parts = item.split("**")
            if len(parts) % 2 == 0:  # asteriscos sin cerrar: el texto tal cual
                parts = [item]
            for index, part in enumerate(parts):
                if part:
                    paragraph.add_run(part).bold = index % 2 == 1

    def note(self, text: str, bold_prefix: str | None = None) -> None:
        """Párrafo sombreado para el estado del informe y las advertencias."""
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        _shade(cell, NOTE_FILL)
        paragraph = cell.paragraphs[0]
        if bold_prefix:
            paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text)
        self.doc.add_paragraph()

    def table(self, header: list[str], rows: list[list[str]], caption: str, widths_cm: list[float] | None = None,
              font_size: float = 8.5, align_numbers: bool = True) -> None:
        self.table_number += 1
        caption_paragraph = self.doc.add_paragraph()
        caption_paragraph.add_run(f"Tabla {self.table_number}. ").bold = True
        caption_paragraph.add_run(caption)
        caption_paragraph.paragraph_format.keep_with_next = True
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = self.text_width_cm
        widths = widths_cm or [total / len(header)] * len(header)
        for row_index, values in enumerate([header] + rows):
            row = table.rows[row_index]
            if row_index == 0:
                _repeat_header(row)
            for column, value in enumerate(values):
                cell = row.cells[column]
                cell.width = Cm(widths[column])
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(str(value))
                run.font.size = Pt(font_size)
                if row_index == 0:
                    run.bold = True
                    _shade(cell, HEADER_FILL)
                elif align_numbers and column > 0 and any(ch.isdigit() for ch in str(value)[:2]):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.doc.add_paragraph()

    def figure(self, path: Path, caption: str, width_cm: float | None = None) -> None:
        self.figure_number += 1
        self.doc.add_picture(str(path), width=Cm(width_cm or min(16.0, self.text_width_cm)))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        paragraph = self.doc.add_paragraph()
        paragraph.add_run(f"Figura {self.figure_number}. ").bold = True
        paragraph.add_run(caption)
        paragraph.runs[-1].font.size = Pt(9)

    def landscape_section(self) -> None:
        section = self.doc.add_section()
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(27.94), Cm(21.59)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.0))

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))


# --- Contenido -------------------------------------------------------------


def _environment_rows(runs: list[Run]) -> list[list[str]]:
    def field(run: Run, getter):
        if not run.available and not run.info:
            return run.status() if getter is specs[-1][1] else "—"
        try:
            value = getter(run)
        except (KeyError, TypeError, AttributeError):
            value = None
        return "—" if value in (None, "", {}) else str(value)

    specs = [
        ("Dónde corre", lambda r: {PC: "PC del equipo, Docker Desktop (WSL 2)", RENDER: "Render, cron job (región Oregon)",
                                   RENDER_C2: "Render, otro cron job (región Oregon)"}[r.label]),
        ("Videos a la vez (--concurrency)", lambda r: (r.info.get("settings") or {}).get("concurrency") or r.first("concurrency")),
        ("Procesador", lambda r: r.info.get("cpu_model") or r.first("cpu_model")),
        ("Límite de CPU (cgroup)", lambda r: num(float(r.info.get("cpu_limit_cores") or r.first("cpu_limit_cores")), 0) + " núcleos"),
        ("Límite de RAM (cgroup)", lambda r: num(float(r.info.get("memory_limit_mb") or r.first("memory_limit_mb")), 0) + " MB"),
        ("CPU visibles (os.cpu_count)", lambda r: r.info.get("host_cpu_count") or r.first("host_cpu_count")),
        ("Hilos por modelo (--threads auto)", lambda r: r.first("threads_requested")),
        ("Sistema", lambda r: r.info.get("platform")),
        ("Python / mediapipe / cv2", lambda r: f"{r.first('python_version')} / {r.first('mediapipe_version')} / {r.first('cv2_version')}"),
        ("torch / ultralytics", lambda r: f"{r.first('torch_version')} / {r.first('ultralytics_version')}"),
        ("Versiones iguales al servicio", lambda r: "sí" if str(r.first("versions_match_service")).lower() == "true" else "NO"),
        ("Corrida (run_id)", lambda r: r.info.get("run_id") or r.first("run_id")),
        ("Inicio (UTC)", lambda r: r.info.get("timestamp_utc")),
        ("Estado", lambda r: r.status()),
    ]
    return [[name] + [field(run, getter) for run in runs] for name, getter in specs]


def _video_rows(runs: list[Run]) -> list[list[str]]:
    source = next((run for run in runs if run.available), None)
    if not source:
        return []
    joints = {v["id"]: v.get("joint") for v in source.info.get("videos", [])}
    seen = {}
    for row in source.rows:
        seen.setdefault(row["video_id"], row)
    order = [v["id"] for v in source.info.get("videos", [])] or list(seen)
    rows = []
    for video_id in order:
        row = seen.get(video_id)
        if not row:
            continue
        rows.append([
            video_id, joints.get(video_id) or "—", row.get("video_frames_reported", ""),
            num(float(row["video_fps"]), 1), num(float(row["video_duration_s"]), 1) + " s",
            num(float(row["video_size_mb"]), 1) + " MB",
            f"{row['video_stored_resolution']} → {row['video_decoded_resolution']}",
            f"{row['video_rotation_meta_deg']}°",
        ])
    return rows


def worker_ram_by_model(run: Run | None) -> dict[str, float]:
    """
    RAM de los procesos que corrieron A LA VEZ: suma de sus picos en cada
    repetición, y el máximo por modelo. Es una cota superior (los dos picos
    pueden no coincidir en el tiempo), que es lo que importa para el límite.
    """
    if run is None:
        return {}
    totals: dict[tuple, float] = {}
    for row in run.rows:
        if row.get("phase") == "measured" and row.get("ram_mb_peak"):
            key = (row["model"], row["video_id"], row.get("trial"), row.get("repetition"))
            totals[key] = totals.get(key, 0.0) + float(row["ram_mb_peak"])
    result: dict[str, float] = {}
    for (model, *_), total in totals.items():
        result[model] = max(result.get(model, 0.0), total)
    return result


def build(pc: Run, render: Run, out_path: Path, image_tag: str | None, render_c2: Run | None = None,
          notes: list[str] | None = None) -> Path:
    runs = [pc, render] + ([render_c2] if render_c2 else [])
    available = [run for run in runs if run.available]
    figures_dir = out_path.parent / (out_path.stem + "_figuras")
    figures_dir.mkdir(parents=True, exist_ok=True)

    summary = summarize(available)
    models_frame = by_model(summary) if not summary.empty else pd.DataFrame()
    pairs = paired(summary)
    decode = decode_by_source(available)
    video_order = [v["id"] for run in available for v in run.info.get("videos", [])]
    video_order = list(dict.fromkeys(video_order))
    both = not pairs.empty
    final = all(run.complete() for run in runs)

    report = Report()
    doc = report.doc

    # --- Portada ---
    title = doc.add_paragraph()
    title_run = title.add_run("Costo computacional de los modelos de pose")
    title_run.bold = True
    title_run.font.size = Pt(20)
    subtitle = doc.add_paragraph()
    subtitle.add_run("Render (instancia 2c-8g) frente al PC del equipo (Docker con 2 CPU / 8 GB)").font.size = Pt(13)
    report.para("Trabajo de grado — Ingeniería de Sistemas, Universidad Antonio Nariño. Grupo LACSER/Bioingeniería.",
                size=9.5)
    report.para(f"Generado: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC por "
                "scripts/build_report.py del repositorio pose-benchmark" + (f" (imagen {image_tag})" if image_tag else "")
                + ".", size=9.5, italic=True)
    report.note(
        " " + " ".join(f"{run.label}: {run.status()}." for run in runs)
        + ("" if final else " Este informe se regenera a medida que llegan datos; las cifras de una corrida en "
           "curso solo incluyen las sesiones modelo-video ya completas."),
        bold_prefix="Estado del informe:" if not final else "Estado del informe: final.",
    )

    # --- 1. Resumen ---
    report.heading("1. Resumen")
    findings = []
    if not models_frame.empty:
        decisive = RENDER if (models_frame["source"] == RENDER).any() else PC
        for model in ordered_models(models_frame["model"]):
            row = models_frame[(models_frame["source"] == decisive) & (models_frame["model"] == model)]
            if row.empty:
                continue
            row = row.iloc[0]
            text = (f"**{model_label(model)}** ({decisive}): {num(row['ms_per_frame'], 0)} ms por frame "
                    f"(×{num(row['realtime_factor'], 2)} la duración del video), {num(row['cores'], 2)} núcleos, "
                    f"RAM pico {num(row['ram_peak'], 0)} MB. Peor job: {num(row['job_s_max'], 1)} s → tiempo "
                    f"{VERDICT_LABELS.get(row['verdict_time'], 'n/d')}, RAM {VERDICT_LABELS.get(row['verdict_ram'], 'n/d')}.")
            if both:
                ratios = pairs[pairs["model"] == model]["ratio_ms"]
                if not ratios.empty:
                    text += f" Frente al PC: {pct_change(geometric_mean(ratios))} de tiempo por frame."
            findings.append(text)
    if both:
        overall = geometric_mean(pairs["ratio_ms"])
        findings.append(f"**En conjunto**, Render tarda {pct_change(overall)} por frame respecto al PC "
                        f"(media geométrica sobre {len(pairs)} combinaciones modelo-video medidas en los dos). "
                        "La RAM y los núcleos usados prácticamente no cambian: lo que cambia es la velocidad de cada núcleo.")
    elif available:
        findings.append("Todavía no hay combinaciones modelo-video completas en los dos entornos: la comparación "
                        "Render-PC aparece en cuanto Render termine y suba sus resultados.")
    c2_frame = models_frame[models_frame["source"] == RENDER_C2] if not models_frame.empty else pd.DataFrame()
    one_frame = models_frame[models_frame["source"] == RENDER] if not models_frame.empty else pd.DataFrame()
    if not c2_frame.empty and not one_frame.empty:
        gains = {m: c2_frame.set_index("model")["throughput"][m] / one_frame.set_index("model")["throughput"][m]
                 for m in ordered_models(c2_frame["model"]) if m in set(one_frame["model"])}
        if gains:
            worst_c2 = c2_frame["job_s_max"].max()
            findings.append(
                "**Con dos videos a la vez (como el worker)**, el worker procesa "
                + ", ".join(f"{pct_change(g)} videos por hora con {model_label(m)}" for m, g in gains.items())
                + f" respecto a uno a la vez. El peor job sube a {num(worst_c2, 1)} s (límite: "
                f"{config.SERVICE_MAX_SHUTDOWN_DELAY_S} s). Detalle en la sección 5.")
    report.bullets(findings or ["Sin datos todavía."])

    # --- 2. Qué se midió y cómo ---
    report.heading("2. Qué se midió y cómo")
    report.heading("2.1 Entornos", 2)
    report.para("Las dos corridas usan la MISMA imagen de contenedor, las mismas librerías y los mismos ocho videos "
                "(verificados por SHA-256). Lo único que cambia es el procesador físico y el sistema que lo comparte.")
    width = report.text_width_cm
    report.table([""] + [run.label for run in runs], _environment_rows(runs),
                 "Condiciones de cada corrida (de run_info.json y de las filas).",
                 widths_cm=[width * 0.25] + [width * 0.75 / len(runs)] * len(runs), align_numbers=False)
    report.heading("2.2 Videos", 2)
    video_rows = _video_rows(runs)
    if video_rows:
        report.table(["Video", "Articulación", "Frames", "fps", "Duración", "Tamaño", "Resolución (guardada → decodificada)", "Rotación"],
                     video_rows, "Videos de TOMA FRONTAL medidos. Todos son de celular en vertical: el metadato de "
                     "rotación se aplica al decodificar y su costo queda dentro de la medición.",
                     widths_cm=[1.5, 2.3, 1.4, 1.1, 1.7, 1.8, 4.3, 1.5])
    report.heading("2.3 Protocolo", 2)
    report.bullets([
        "Cuatro modelos: YOLO26-nano, MediaPipe Lite, Full y Heavy (Full es referencia, no candidato).",
        "Cada combinación modelo-video corre en un proceso propio: 1 corrida de calentamiento (no entra al análisis) y "
        "3 repeticiones medidas. Cada repetición construye el modelo desde cero, como un job del servicio.",
        "Una repetición procesa el video completo: decodificar, aplicar la rotación, convertir de color e inferir, "
        "frame a frame. job = construir el modelo + procesar el video. Como el calentamiento ya importó las "
        "librerías, «construir el modelo» en las repeticiones medidas es solo crear el objeto (centésimas de "
        "segundo); el arranque en frío del proceso queda en la fila de calentamiento.",
        "CPU: segundos de CPU del proceso ÷ segundos de reloj (núcleos usados). RAM: pico de memoria residente del "
        "proceso (VmHWM). Además, contadores del cgroup (tiempo frenado por la cuota, OOM-kills).",
        "Concurrencia 1 (un video a la vez) y hilos según la cuota de CPU (2). El worker del servicio corre 2 videos "
        "a la vez: eso es otra corrida (ver Limitaciones).",
        "Sin red durante la medición en el PC; en Render los videos se bajan de Google Drive antes de medir, fuera "
        "de los tiempos.",
    ])
    report.heading("2.4 Criterios de aceptación", 2)
    share = config.BUDGET_RAM_MB * stats.RAM_SAFE_FRACTION
    report.bullets([
        f"**Tiempo (peor job medido):** aceptable si ≤ {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s (lo que Render espera "
        f"a un job al redesplegar el worker); condicionado si ≤ {config.SERVICE_STUCK_JOB_TIMEOUT_S} s (el barrido de "
        "jobs atascados del servicio); no aceptable por encima.",
        f"**RAM:** pico × {config.SERVICE_WORKER_CONCURRENCY} procesos del worker ≤ {num(share, 0)} MB (80 % de "
        f"{config.BUDGET_RAM_MB} MB) es aceptable; ≤ {config.BUDGET_RAM_MB} MB, condicionado.",
        f"**Validez de la medición:** al menos {stats.MIN_REPETITIONS} repeticiones, coeficiente de variación ≤ "
        f"{num(stats.MAX_CV_PCT, 0)} % y detección ≥ {num(stats.MIN_DETECTION_RATE * 100, 0)} %.",
        "Detalle y referencias en docs/criterios_rendimiento.md del repositorio.",
    ])

    # --- 3. Resultados por modelo ---
    report.heading("3. Resultados por modelo")
    if models_frame.empty:
        report.para("Sin sesiones completas todavía.")
    else:
        header = ["Modelo", "Entorno", "ms/frame (mediana de videos)", "Rango entre videos", "× duración del video",
                  "Núcleos", "RAM pico", "Peor job", "CV máx.", "Veredicto tiempo / RAM"]
        rows = []
        for model in ordered_models(models_frame["model"]):
            for source in (PC, RENDER):
                match = models_frame[(models_frame["source"] == source) & (models_frame["model"] == model)]
                if match.empty:
                    continue
                r = match.iloc[0]
                rows.append([model_label(model), f"{source} ({r['videos']} videos)", num(r["ms_per_frame"], 1),
                             f"{num(r['ms_per_frame_min'], 0)}–{num(r['ms_per_frame_max'], 0)}",
                             "×" + num(r["realtime_factor"], 2), num(r["cores"], 2), num(r["ram_peak"], 0) + " MB",
                             num(r["job_s_max"], 1) + " s", num(r["cv_max"], 1) + " %",
                             f"{VERDICT_LABELS.get(r['verdict_time'], 'n/d')} / {VERDICT_LABELS.get(r['verdict_ram'], 'n/d')}"
                             + ("" if r["valid_all"] else " (*)")])
        report.table(header, rows, "Resultados agregados sobre los videos. (*) alguna sesión no cumple los criterios "
                     "de validez (ver Anexo A).", widths_cm=[2.1, 1.9, 1.6, 1.5, 1.5, 1.2, 1.5, 1.4, 1.2, 2.3], font_size=8)
        videos_measured = summary["video_id"].nunique()
        path = source_bars(models_frame, "ms_per_frame", f"Tiempo por frame (mediana de {videos_measured} videos)", "ms por frame",
                           figures_dir / "01_ms_por_frame.png", SOURCE_COLORS, spread=("ms_per_frame_min", "ms_per_frame_max"))
        report.figure(path, "Tiempo por frame de cada modelo. La barra es la mediana entre videos; el bigote va del "
                      "video más rápido al más lento. Incluye decodificar y rotar cada frame.")
        if both:
            path = ratio_dots(pairs, figures_dir / "02_render_vs_pc.png")
            report.figure(path, "Cuántas veces más lento es Render que el PC, por modelo. Cada punto es un video; "
                          "la raya negra es la media geométrica. A la derecha de la línea punteada, Render es más lento.")

    # --- 4. Tiempo por video y límites del servicio ---
    report.heading("4. Tiempo por job y límites del servicio")
    decisive = RENDER if render.available and (summary["environment"] == RENDER).any() else PC
    if not summary.empty:
        subset = summary[summary["environment"] == decisive]
        path = figures_dir / f"03_job_max_{decisive.lower()}.png"
        video_bars(subset, "job_s_max", video_order, f"Peor job por video — {decisive}", "segundos", path,
                   (config.SERVICE_MAX_SHUTDOWN_DELAY_S, f"apagado del worker: {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s"))
        report.figure(path, f"Peor job medido (construir el modelo + procesar el video) en cada video, en {decisive}. "
                      f"La línea es el límite de {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s para que un job no se corte "
                      "al redesplegar el worker.")
        rows = []
        for model in ordered_models(models_frame["model"]):
            match = models_frame[(models_frame["source"] == decisive) & (models_frame["model"] == model)]
            if match.empty:
                continue
            r = match.iloc[0]
            rows.append([model_label(model), "×" + num(r["realtime_factor"], 2), num(r["model_init_s"], 2) + " s",
                         num(r["max_video_s"], 0) + " s (" + num(r["max_video_s"] / 60, 1) + " min)",
                         num(r["throughput"], 0)])
        report.table(["Modelo", "× duración del video", "Construir el modelo", "Video más largo que cabe en 300 s",
                      "Videos por hora (1 a la vez)"], rows,
                     f"Lo que implica el ritmo medido en {decisive} para el servicio. El costo escala con la duración "
                     "del video (frames), no con su tamaño en MB: el límite de 200 MB del servicio no acota la duración.",
                     font_size=8.5)

    # --- 5. Dos videos a la vez ---
    report.heading("5. Render con dos videos a la vez (como el worker)")
    if c2_frame.empty:
        report.para("Pendiente: la corrida de Render con --concurrency 2 todavía no tiene sesiones completas.")
    else:
        report.para(f"El worker del servicio corre con --concurrency={config.SERVICE_WORKER_CONCURRENCY}: dos videos a "
                    "la vez en las mismas 2 CPU. Esta corrida repite la suite así, en otra instancia 2c-8g de Render, "
                    "con 1 hilo por proceso (--threads auto reparte los 2 núcleos entre los 2 procesos). Se compara "
                    "con la corrida de Render de un video a la vez; el PC no entra en esta sección.")
        worker_ram = worker_ram_by_model(render_c2)
        rows = []
        for model in ordered_models(c2_frame["model"]):
            two = c2_frame[c2_frame["model"] == model].iloc[0]
            one = one_frame[one_frame["model"] == model]
            one = one.iloc[0] if not one.empty else None
            gain = two["throughput"] / one["throughput"] if one is not None else None
            rows.append([model_label(model),
                         num(one["ms_per_frame"], 1) if one is not None else "n/d", num(two["ms_per_frame"], 1),
                         num(one["job_s_max"], 1) + " s" if one is not None else "n/d", num(two["job_s_max"], 1) + " s",
                         num(one["throughput"], 0) if one is not None else "n/d", num(two["throughput"], 0),
                         pct_change(gain), num(worker_ram.get(model), 0) + " MB",
                         num(two["max_video_s"], 0) + " s", VERDICT_LABELS.get(two["verdict_time"], "n/d")])
        report.table(["Modelo", "ms/frame, 1 a la vez", "ms/frame, 2 a la vez (cada proceso)", "Peor job, 1",
                      "Peor job, 2", "Videos/hora, 1", "Videos/hora, 2", "Cambio en videos/hora",
                      "RAM de los 2 procesos", "Video más largo en 300 s (2 a la vez)", "Veredicto tiempo (2)"], rows,
                     "Render con uno y con dos videos a la vez. Videos por hora = videos que termina el worker completo "
                     "(con 2 a la vez, dos jobs en paralelo). RAM de los 2 procesos: suma de sus picos, una cota superior.",
                     widths_cm=[2.0, 1.4, 1.6, 1.3, 1.3, 1.4, 1.4, 1.5, 1.5, 1.6, 1.6], font_size=7.5)
        compare = models_frame[models_frame["source"].isin([RENDER, RENDER_C2])]
        path = source_bars(compare, "throughput", "Videos por hora que termina el worker", "videos por hora",
                           figures_dir / "05a_videos_por_hora.png", CONCURRENCY_COLORS)
        report.figure(path, f"Videos por hora en Render: «{RENDER}» procesa uno a la vez; «{RENDER_C2}», dos a la vez "
                      "(como el worker del servicio). Mediana de los 8 videos, cada uno de 11 a 17 s.")
        path = source_bars(compare, "ms_per_frame", "Tiempo por frame de cada proceso", "ms por frame",
                           figures_dir / "05b_ms_por_frame_concurrencia.png", CONCURRENCY_COLORS,
                           spread=("ms_per_frame_min", "ms_per_frame_max"))
        report.figure(path, "Cuánto se alarga cada video cuando comparte las 2 CPU con otro.")
        subset = summary[summary["environment"] == RENDER_C2]
        path = video_bars(subset, "job_s_max", video_order, f"Peor job por video — {RENDER_C2}", "segundos",
                          figures_dir / "05c_job_max_render_x2.png",
                          (config.SERVICE_MAX_SHUTDOWN_DELAY_S, f"apagado del worker: {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s"))
        report.figure(path, "Peor job de cada video con dos videos a la vez: es el tiempo que vería el servicio.")

    # --- 6. CPU ---
    report.heading("6. Uso de CPU")
    if not models_frame.empty:
        path = source_bars(models_frame, "cores", "Núcleos usados (mediana de videos)", "núcleos",
                           figures_dir / "04_nucleos.png", SOURCE_COLORS, decimals=2,
                           reference=(config.BUDGET_CPU_CORES, f"presupuesto {config.RENDER_PLAN}: {config.BUDGET_CPU_CORES} núcleos"))
        report.figure(path, "Núcleos usados en promedio durante el procesamiento (segundos de CPU ÷ segundos de reloj).")
        items = []
        for model in ordered_models(models_frame["model"]):
            values = {s: models_frame[(models_frame["source"] == s) & (models_frame["model"] == model)] for s in (PC, RENDER)}
            parts = [f"{s}: {num(v.iloc[0]['cores'], 2)} núcleos, frenado por la cuota en el "
                     f"{num(v.iloc[0]['throttled_pct'], 0)} % de los periodos" for s, v in values.items() if not v.empty]
            items.append(f"**{model_label(model)}** — " + "; ".join(parts) + ".")
        report.bullets(items)
        report.para("«Frenado por la cuota» es la fracción de periodos de 100 ms del planificador en que el contenedor "
                    "agotó sus 2 núcleos y tuvo que esperar (cpu.stat del cgroup). Un modelo que usa ~2 núcleos vive "
                    "pegado al límite; uno que usa ~1 casi no lo toca.", size=9.5)

    # --- 7. Memoria ---
    report.heading("7. Memoria")
    if not models_frame.empty:
        path = source_bars(models_frame, "ram_peak", "RAM pico por proceso (peor video)", "MB", figures_dir / "05_ram.png", SOURCE_COLORS)
        per_process = share / config.SERVICE_WORKER_CONCURRENCY
        worst = models_frame["ram_peak"].max()
        report.figure(path, f"RAM pico de un proceso de medición. El margen por proceso con {config.SERVICE_WORKER_CONCURRENCY} "
                      f"en paralelo es {num(per_process, 0)} MB; el peor pico medido es {num(worst, 0)} MB "
                      f"({num(worst / per_process * 100, 0)} % del margen).")
        drift = [f"**{model_label(m)}**: " + "; ".join(
                    f"{s} +{num(r.iloc[0]['ram_drift'], 0)} MB" for s in (PC, RENDER)
                    if not (r := models_frame[(models_frame['source'] == s) & (models_frame['model'] == m)]).empty)
                 for m in ordered_models(models_frame["model"])]
        report.para("Deriva de memoria: cuánto subió la RAM de partida entre la primera y la última repetición de un "
                    "mismo proceso (el modelo se crea de nuevo en cada repetición, como hace el worker del servicio). "
                    "Es la fragmentación de glibc documentada en docs/criterios_rendimiento.md:")
        report.bullets(drift)

    # --- 8. Diferencias ---
    report.heading("8. Diferencias entre Render y el PC")
    if both:
        items = []
        per_model = {m: geometric_mean(pairs[pairs["model"] == m]["ratio_ms"]) for m in ordered_models(pairs["model"])}
        slowest = max(per_model, key=per_model.get)
        fastest = min(per_model, key=per_model.get)
        items.append(f"**Velocidad por núcleo.** Render es {pct_change(geometric_mean(pairs['ratio_ms']))} en tiempo "
                     f"por frame. El efecto no es igual en todos los modelos: va de {pct_change(per_model[fastest])} "
                     f"({model_label(fastest)}) a {pct_change(per_model[slowest])} ({model_label(slowest)}). Los dos "
                     "procesadores son de la misma generación (Zen 3); el EPYC 7R13 de Render es un procesador de "
                     "servidor con menor frecuencia por núcleo que el Ryzen 5 5600G del PC.")
        if PC in decode and RENDER in decode:
            items.append(f"**Decodificar y rotar un frame** (sin modelo): {num(decode[PC], 1)} ms en el PC y "
                         f"{num(decode[RENDER], 1)} ms en Render ({pct_change(decode[RENDER] / decode[PC])}).")
        ram_diff = pairs["ram_diff_mb"]
        items.append(f"**RAM.** La diferencia de pico entre entornos va de {num(ram_diff.min(), 0)} a "
                     f"{num(ram_diff.max(), 0)} MB (Render − PC): la memoria depende del modelo y del video, no del "
                     "procesador. Lo medido en el PC vale para Render.")
        cores_diff = pairs["cores_diff"]
        items.append(f"**Núcleos usados.** Render − PC entre {num(cores_diff.min(), 2)} y {num(cores_diff.max(), 2)} "
                     "núcleos: cada modelo ocupa la misma fracción del presupuesto en los dos entornos.")
        cv_cols = [c for c in ("processing_s_cv_pct_pc", "processing_s_cv_pct_render") if c in pairs.columns]
        if len(cv_cols) == 2:
            items.append(f"**Estabilidad.** Coeficiente de variación mediano entre repeticiones: "
                         f"{num(pairs[cv_cols[0]].median(), 1)} % en el PC y {num(pairs[cv_cols[1]].median(), 1)} % "
                         "en Render.")
        if {"cgroup_throttled_pct_median_pc", "cgroup_throttled_pct_median_render"} <= set(pairs.columns):
            items.append(f"**Cuota de CPU.** Periodos frenados (mediana): {num(pairs['cgroup_throttled_pct_median_pc'].median(), 0)} % "
                         f"en el PC y {num(pairs['cgroup_throttled_pct_median_render'].median(), 0)} % en Render. Los "
                         "dos limitan por cuota del cgroup v2, pero el planificador de WSL 2 y el de la instancia de "
                         "Render reparten los periodos de forma distinta; los núcleos usados, que es lo que importa, coinciden.")
        report.bullets(items)
        report.table(
            ["Modelo", "Video", "ms/frame PC", "ms/frame Render", "Render ÷ PC", "RAM pico PC", "RAM pico Render", "Núcleos PC", "Núcleos Render"],
            [[model_label(r["model"]), r["video_id"], num(r["ms_per_frame_median_pc"], 1), num(r["ms_per_frame_median_render"], 1),
              "×" + num(r["ratio_ms"], 2), num(r["ram_mb_peak_max_pc"], 0), num(r["ram_mb_peak_max_render"], 0),
              num(r["cpu_cores_used_median_pc"], 2), num(r["cpu_cores_used_median_render"], 2)]
             for _, r in pairs.sort_values(["model", "video_id"], key=lambda col: col.map(
                 {m: i for i, m in enumerate(MODEL_COLORS)}) if col.name == "model" else col).iterrows()],
            "Comparación pareada por modelo y video (medianas de las 3 repeticiones; RAM: máximo).",
            widths_cm=[2.3, 1.3, 1.6, 1.8, 1.6, 1.6, 1.8, 1.5, 1.6], font_size=8)
    else:
        report.para("Pendiente: hace falta que las dos corridas tengan al menos una combinación modelo-video completa.")

    # --- 9. Validez y limitaciones ---
    report.heading("9. Validez y limitaciones")
    limitations = [
        "**Una sola instancia de Render y una sola corrida.** La nube varía entre instancias y a lo largo del día "
        "(Laaber et al., 2019; Leitner & Cito, 2016). Para afirmar un tiempo de Render con confianza conviene repetir "
        "la suite en otros momentos (--trials con orden barajado, metodología RMIT de Abedi & Brecht, 2017).",
        ("**Dos videos a la vez, una sola corrida.** La sección 5 mide el caso del worker (--concurrency 2) en una "
         "instancia de Render; el PC no se midió así." if not c2_frame.empty else
         "**Concurrencia 1.** El worker del servicio procesa 2 videos a la vez en las mismas 2 CPU; con dos jobs "
         "simultáneos el tiempo por video sube. Falta la corrida con --concurrency 2 en Render."),
        "**El PC no es una máquina dedicada.** Docker Desktop comparte el procesador con Windows; durante la "
        "corrida local también se generó este informe. La VM de WSL 2 tiene ~7,8 GB, por debajo del límite nominal "
        "de 8 GB (irrelevante aquí: ningún modelo pasa de 1 GB).",
        "**Videos de 11 a 17 s.** Los tiempos por job escalan con la duración; la tabla de la sección 4 da el video "
        "más largo que cabe en cada límite al ritmo medido.",
        "**Mide el costo del modelo, no el pipeline del servicio.** La validación, el cálculo de ángulos y la subida "
        "de resultados del servicio no están incluidos (eso lo mide app/tools/benchmark_pipeline.py del servicio).",
    ]
    if not summary.empty and not summary["measurement_valid"].all():
        invalid = summary[~summary["measurement_valid"]]
        limitations.insert(0, f"**{len(invalid)} sesiones no cumplen los criterios de validez** (detalle en el Anexo A): "
                           + "; ".join(f"{r['environment']} {model_label(r['model'])} {r['video_id']}: "
                                       + re.sub(r"(\d)\.(\d)", r"\1,\2", r["warnings"].split(";")[0])
                                       for _, r in invalid.head(6).iterrows()) + ".")
    # Notas del equipo sobre lo que pasó durante las corridas (--note): los
    # datos no se tocan; lo que haya que saber para leerlos se dice acá.
    limitations[0:0] = [f"**Nota.** {note}" for note in notes or []]
    report.bullets(limitations)

    # --- Anexos (horizontal) ---
    if not summary.empty:
        report.landscape_section()
        report.heading("Anexo A. Resumen completo por entorno, modelo y video")
        rows = []
        model_rank = {m: i for i, m in enumerate(ordered_models(summary["model"]))}
        video_rank = {v: i for i, v in enumerate(video_order)}
        ordered = summary.assign(_m=summary["model"].map(model_rank), _v=summary["video_id"].map(video_rank).fillna(99))
        for _, r in ordered.sort_values(["_m", "_v", "environment"]).iterrows():
            rows.append([r["environment"], model_label(r["model"]), r["video_id"], str(r["n"]),
                         num(r["processing_s_median"], 2), num(r["processing_s_cv_pct"], 1),
                         f"{num(r['processing_s_ci95_low'], 2)}–{num(r['processing_s_ci95_high'], 2)}",
                         num(r["ms_per_frame_median"], 1), num(r["infer_ms_p90_median"], 1),
                         num(r["model_init_s_median"], 2), num(r["job_s_max"], 2), num(r["cpu_cores_used_median"], 2),
                         num(r["ram_mb_peak_max"], 0), num(r["ram_drift_mb"], 0), num(r["detection_rate_min"] * 100, 0) + " %",
                         VERDICT_LABELS.get(r["verdict_time"], "n/d"), "sí" if r["measurement_valid"] else "no"])
        report.table(["Entorno", "Modelo", "Video", "n", "Proceso (s, mediana)", "CV %", "IC 95 % (s)", "ms/frame",
                      "Inferencia p90 (ms)", "Init (s)", "Peor job (s)", "Núcleos", "RAM pico (MB)", "Deriva (MB)",
                      "Detección mín.", "Veredicto tiempo", "Válida"], rows,
                     "Una fila por entorno, modelo y video (3 repeticiones medidas; sin calentamiento).",
                     widths_cm=[1.8, 2.2, 1.1, 0.6, 1.5, 1.0, 2.1, 1.3, 1.5, 1.0, 1.4, 1.3, 1.5, 1.3, 1.5, 1.8, 1.1],
                     font_size=7)
        report.heading("Anexo B. Tiempo por frame en cada video")
        for index, source in enumerate(run.label for run in runs):
            subset = summary[summary["environment"] == source]
            if subset.empty:
                continue
            path = figures_dir / f"09_ms_por_video_{index + 1}.png"
            video_bars(subset, "ms_per_frame_median", video_order, f"Tiempo por frame por video — {source}",
                       "ms por frame (mediana)", path)
            report.figure(path, f"Tiempo por frame de cada modelo en cada video, {source}.", width_cm=21)

    report.save(out_path)
    print(f"  -> {out_path}")
    return out_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m scripts.build_report", description=__doc__.split("\n\n")[0])
    parser.add_argument("--pc", type=Path, help="carpeta de la corrida en el PC (results/<run_id>)")
    parser.add_argument("--render", type=Path, help="carpeta de la corrida en Render (results/<run_id>)")
    parser.add_argument("--out", type=Path, required=True, help="ruta del .docx")
    parser.add_argument("--render-c2", type=Path, help="carpeta de la corrida en Render con --concurrency 2")
    parser.add_argument("--image-tag", default=None, help="etiqueta de la imagen medida (p. ej. davidgr05/mvcbenchmark:5a87fa0)")
    parser.add_argument("--note", action="append", default=[],
                        help="nota sobre las corridas para la sección de validez (se puede repetir)")
    args = parser.parse_args(argv)
    if not args.pc and not args.render:
        parser.error("hace falta --pc, --render o los dos")
    build(Run(PC, args.pc), Run(RENDER, args.render), args.out, args.image_tag,
          Run(RENDER_C2, args.render_c2) if args.render_c2 else None, args.note)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
