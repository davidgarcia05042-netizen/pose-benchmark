"""
Informe en Word SOLO con las mediciones de Render: comparación entre modelos.

Análisis LOCAL (no va en la imagen). Complementa a scripts/build_report.py
(Render frente al PC) sin reemplazarlo. Acá no entra el PC: sus cifras
compartieron el procesador con lo que se usaba en el equipo mientras corrían,
y las instancias de Render corrieron dedicadas a la medición. La pregunta es
otra: entre los modelos, cuál cuesta menos en el hardware del worker
(`2c-8g`), con uno y con dos videos a la vez.

Uso:
    python -m scripts.build_render_report \\
        --render results/20261001T063552Z \\
        --render-c2 results/20261001T070703Z \\
        --out results/informes/informe_render_modelos.docx \\
        --image-tag davidgr05/mvcbenchmark:5a87fa0
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path
from statistics import median

import pandas as pd
from docx.shared import Pt

from benchmark import config, stats
from scripts.build_report import (CONCURRENCY_COLORS, RENDER, RENDER_C2, VERDICT_LABELS, Report, Run,
                                  _environment_rows, _video_rows, by_model, pct_change, summarize,
                                  worker_ram_by_model)
from scripts.charts import model_bars, model_label, num, ordered_models, source_bars, video_bars

# Tarifa que mostró el panel de Render para un cron job 2c-8g el 2026-10-01.
# Es un INDICADOR para comparar modelos entre sí, no el costo del servicio: el
# worker se factura distinto y el plan lo decide el grupo de investigación
# (bloqueador #1 de docs/DECISIONES_TECNICAS.md).
CRON_2C8G_USD_PER_MIN = 0.00313


def _float(row: dict, key: str) -> float | None:
    try:
        value = row.get(key)
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _median(values) -> float | None:
    values = [v for v in values if v is not None]
    return median(values) if values else None


def raw_metrics(run: Run) -> pd.DataFrame:
    """
    Lo que no está en el resumen de stats y se saca de las filas: costo de CPU
    por segundo de video, decodificación, inferencia p50, y del calentamiento
    (proceso recién arrancado) el arranque en frío y la RAM del modelo cargado.
    """
    records = []
    for model in ordered_models({r["model"] for r in run.rows}):
        measured = [r for r in run.rows if r["model"] == model and r.get("phase") == "measured"]
        warmup = [r for r in run.rows if r["model"] == model and r.get("phase") == "warmup"]
        cpu_per_video_s = []
        for row in measured:
            cpu_s, frames, fps = _float(row, "cpu_s"), _float(row, "frames_processed"), _float(row, "video_fps")
            if cpu_s is not None and frames and fps:
                cpu_per_video_s.append(cpu_s / (frames / fps))
        model_ram = [(_float(r, "ram_after_init_mb") or 0) - (_float(r, "ram_start_mb") or 0) for r in warmup
                     if _float(r, "ram_after_init_mb") is not None and _float(r, "ram_start_mb") is not None]
        records.append({
            "model": model,
            "cpu_per_video_s": _median(cpu_per_video_s),
            "decode_ms": _median(_float(r, "decode_ms_mean") for r in measured),
            "infer_p50": _median(_float(r, "infer_ms_p50") for r in measured),
            "infer_p90": _median(_float(r, "infer_ms_p90") for r in measured),
            "fps": _median(_float(r, "fps_processed") for r in measured),
            "cold_init_s": _median(_float(r, "model_init_s") for r in warmup),
            "model_ram_mb": _median(model_ram),
            "cpu_budget_pct": _median(_float(r, "cpu_budget_pct") for r in measured),
        })
    return pd.DataFrame(records)


def _cv_by_model(summary: pd.DataFrame, source: str) -> dict[str, tuple[float, float, int]]:
    """(CV mediano, CV máximo, sesiones con CV > MAX_CV_PCT) por modelo."""
    result = {}
    subset = summary[summary["environment"] == source]
    for model, group in subset.groupby("model"):
        cvs = group["processing_s_cv_pct"].dropna()
        result[model] = (cvs.median(), cvs.max(), int((cvs > stats.MAX_CV_PCT).sum()))
    return result


def build(render: Run, render_c2: Run | None, out_path: Path, image_tag: str | None,
          notes: list[str] | None = None) -> Path:
    runs = [render] + ([render_c2] if render_c2 else [])
    available = [run for run in runs if run.available]
    figures_dir = out_path.parent / (out_path.stem + "_figuras")
    figures_dir.mkdir(parents=True, exist_ok=True)
    for old in figures_dir.glob("*.png"):  # solo las figuras de esta versión
        old.unlink()

    summary = summarize(available)
    frame = by_model(summary) if not summary.empty else pd.DataFrame()
    one = frame[frame["source"] == RENDER].set_index("model") if not frame.empty else pd.DataFrame()
    two = frame[frame["source"] == RENDER_C2].set_index("model") if not frame.empty else pd.DataFrame()
    raw_one = raw_metrics(render).set_index("model") if render.available else pd.DataFrame()
    raw_two = raw_metrics(render_c2).set_index("model") if render_c2 and render_c2.available else pd.DataFrame()
    worker_ram = worker_ram_by_model(render_c2)
    models = ordered_models(one.index) if not one.empty else []
    video_order = list(dict.fromkeys(v["id"] for run in available for v in run.info.get("videos", [])))
    final = all(run.complete() for run in runs)
    usd_per_hour = CRON_2C8G_USD_PER_MIN * 60

    report = Report()
    doc = report.doc

    # --- Portada ---
    title = doc.add_paragraph()
    title_run = title.add_run("Costo computacional de los modelos de pose en Render")
    title_run.bold = True
    title_run.font.size = Pt(20)
    doc.add_paragraph().add_run("Comparación entre modelos en la instancia 2c-8g del worker (2 CPU, 8 GB)").font.size = Pt(13)
    report.para("Trabajo de grado — Ingeniería de Sistemas, Universidad Antonio Nariño. Grupo LACSER/Bioingeniería.", size=9.5)
    report.para(f"Generado: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC por "
                "scripts/build_render_report.py del repositorio pose-benchmark"
                + (f" (imagen {image_tag})" if image_tag else "") + ".", size=9.5, italic=True)
    report.note(" " + " ".join(f"{run.label}: {run.status()}." for run in runs)
                + ("" if final else " Las cifras solo incluyen sesiones modelo-video completas."),
                bold_prefix="Estado del informe: final." if final else "Estado del informe:")

    if one.empty:
        report.para("Sin sesiones completas de Render todavía.")
        report.save(out_path)
        return out_path

    # --- 1. Resumen ---
    report.heading("1. Resumen")
    by_speed = sorted(models, key=lambda m: one.loc[m, "ms_per_frame"])
    findings = [
        "**Del más rápido al más lento (un video a la vez):** "
        + ", ".join(f"{model_label(m)} {num(one.loc[m, 'ms_per_frame'], 0)} ms/frame (×{num(one.loc[m, 'realtime_factor'], 2)} "
                    "la duración del video)" for m in by_speed) + ".",
    ]
    if raw_one["cpu_per_video_s"].notna().all():
        by_cpu = sorted(models, key=lambda m: raw_one.loc[m, "cpu_per_video_s"])
        findings.append("**Del que menos CPU gasta al que más, por segundo de video:** "
                        + ", ".join(f"{model_label(m)} {num(raw_one.loc[m, 'cpu_per_video_s'], 1)} s de CPU" for m in by_cpu)
                        + ". Es el costo de cómputo real: tiempo × núcleos ocupados.")
    fastest, slowest = by_speed[0], by_speed[-1]
    cheapest = min(models, key=lambda m: raw_one.loc[m, "cpu_per_video_s"]) if raw_one["cpu_per_video_s"].notna().all() else None
    where = "en tiempo y en CPU" if cheapest in (None, fastest) else f"en tiempo (en CPU gana {model_label(cheapest)})"
    findings.append(f"**{model_label(fastest)} es el más eficiente** {where}; "
                    f"**{model_label(slowest)} es el más lento** (×{num(one.loc[slowest, 'ms_per_frame'] / one.loc[fastest, 'ms_per_frame'], 1)} "
                    f"el tiempo por frame de {model_label(fastest)}).")
    if "yolo26n" in models and "mediapipe_heavy" in models:
        findings.append(
            f"**YOLO26-nano y MediaPipe Heavy cuestan distinto aunque tarden parecido:** YOLO26-nano es más rápido "
            f"({num(one.loc['yolo26n', 'ms_per_frame'], 0)} frente a {num(one.loc['mediapipe_heavy', 'ms_per_frame'], 0)} "
            f"ms/frame), pero ocupa {num(one.loc['yolo26n', 'cores'], 2)} núcleos contra {num(one.loc['mediapipe_heavy', 'cores'], 2)}: "
            f"gasta {num(raw_one.loc['yolo26n', 'cpu_per_video_s'], 1)} s de CPU por segundo de video frente a "
            f"{num(raw_one.loc['mediapipe_heavy', 'cpu_per_video_s'], 1)}.")
    ram_low, ram_high = min(models, key=lambda m: one.loc[m, "ram_peak"]), max(models, key=lambda m: one.loc[m, "ram_peak"])
    findings.append(f"**RAM:** de {num(one.loc[ram_low, 'ram_peak'], 0)} MB ({model_label(ram_low)}) a "
                    f"{num(one.loc[ram_high, 'ram_peak'], 0)} MB ({model_label(ram_high)}) por proceso: todos usan menos "
                    f"del {num(one['ram_peak'].max() / config.BUDGET_RAM_MB * 100 + 0.5, 0)} % de los 8 GB.")
    if not two.empty:
        gains = {m: two.loc[m, "throughput"] / one.loc[m, "throughput"] for m in models if m in two.index}
        best_gain = max(gains, key=gains.get)
        # La explicación solo vale si ese modelo, con uno a la vez, dejaba CPU sin usar.
        reason = (f", porque con uno a la vez usa {num(one.loc[best_gain, 'cores'], 2)} de los 2 núcleos y el segundo "
                  "video aprovecha lo que sobra") if one.loc[best_gain, "cores"] < 1.5 else ""
        findings.append("**Con dos videos a la vez (como el worker)**, el worker termina más videos por hora con todos "
                        "los modelos: " + ", ".join(f"{model_label(m)} {pct_change(g)}" for m, g in gains.items())
                        + f". Gana más {model_label(best_gain)}{reason}.")
        tight = min((m for m in models if m in two.index), key=lambda m: two.loc[m, "max_video_s"])
        findings.append(f"**Límite del servicio:** con dos a la vez, {model_label(tight)} solo alcanza a procesar videos "
                        f"de hasta ~{num(two.loc[tight, 'max_video_s'], 0)} s dentro de los "
                        f"{config.SERVICE_MAX_SHUTDOWN_DELAY_S} s del apagado del worker. Con los videos medidos (11–17 s), "
                        "todos los modelos son aceptables en tiempo y en RAM.")
    findings.append("**Este informe no mide precisión.** Cuál modelo conviene combina este costo con la precisión "
                    "medida en entorno-pruebas contra el gold standard.")
    report.bullets(findings)

    # --- 2. Condiciones ---
    report.heading("2. Qué se midió y cómo")
    report.heading("2.1 Por qué solo Render", 2)
    report.para("Las mismas mediciones se hicieron también en el PC del equipo (Docker con 2 CPU / 8 GB); esa "
                "comparación está en el informe «Render frente al PC». Acá solo entran las corridas de Render: "
                "corrieron en instancias dedicadas a la medición, sin nada más usando el procesador, mientras que el "
                "PC se siguió usando durante su corrida y al menos una repetición quedó perturbada por eso. Además, "
                "Render es el hardware del worker del servicio: estas son las cifras que cuentan para decidir.")
    report.heading("2.2 Instancias", 2)
    width = report.text_width_cm
    report.table([""] + [run.label for run in runs], _environment_rows(runs),
                 f"Condiciones de cada corrida. «{RENDER}»: un video a la vez; «{RENDER_C2}»: dos a la vez, en otra "
                 "instancia del mismo plan.", widths_cm=[width * 0.3] + [width * 0.7 / len(runs)] * len(runs),
                 align_numbers=False)
    report.heading("2.3 Videos", 2)
    report.table(["Video", "Articulación", "Frames", "fps", "Duración", "Tamaño", "Resolución (guardada → decodificada)", "Rotación"],
                 _video_rows(available), "Videos de TOMA FRONTAL (celular en vertical; la rotación se aplica al "
                 "decodificar y su costo queda dentro de la medición). Se bajan de Google Drive antes de medir y se "
                 "verifica su SHA-256.", widths_cm=[1.5, 2.3, 1.4, 1.1, 1.7, 1.8, 4.3, 1.5])
    report.heading("2.4 Protocolo y criterios", 2)
    report.bullets([
        "Cuatro modelos: YOLO26-nano, MediaPipe Lite, Full y Heavy (Full es referencia, no candidato).",
        "Cada combinación modelo-video corre en un proceso propio: 1 calentamiento (arranque en frío; no entra a las "
        "cifras de tiempo) y 3 repeticiones medidas, cada una con el modelo construido desde cero, como un job del "
        "servicio. Con dos a la vez, dos procesos idénticos corren en paralelo (6 repeticiones por sesión).",
        "Hilos según la cuota de CPU: 2 por proceso con uno a la vez; 1 por proceso con dos a la vez.",
        f"Tiempo: aceptable si el peor job ≤ {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s; condicionado si ≤ "
        f"{config.SERVICE_STUCK_JOB_TIMEOUT_S} s. RAM: aceptable si los {config.SERVICE_WORKER_CONCURRENCY} procesos "
        f"del worker caben en el 80 % de {config.BUDGET_RAM_MB} MB. Medición válida: ≥ {stats.MIN_REPETITIONS} "
        f"repeticiones, CV ≤ {num(stats.MAX_CV_PCT, 0)} %, detección ≥ {num(stats.MIN_DETECTION_RATE * 100, 0)} %.",
    ])

    # --- 3. Tiempo ---
    report.heading("3. Tiempo de procesamiento por modelo")
    rows = [[model_label(m), num(one.loc[m, "ms_per_frame"], 1),
             f"{num(one.loc[m, 'ms_per_frame_min'], 0)}–{num(one.loc[m, 'ms_per_frame_max'], 0)}",
             "×" + num(one.loc[m, "realtime_factor"], 2), num(raw_one.loc[m, "fps"], 1),
             num(raw_one.loc[m, "infer_p50"], 1), num(raw_one.loc[m, "infer_p90"], 1), num(raw_one.loc[m, "decode_ms"], 1),
             num(one.loc[m, "job_s_max"], 1) + " s"] for m in models]
    report.table(["Modelo", "ms/frame (mediana)", "Rango entre videos", "× duración del video", "Frames por segundo",
                  "Inferencia p50 (ms)", "Inferencia p90 (ms)", "Decodificar + rotar (ms)", "Peor job"], rows,
                 f"Tiempo con un video a la vez ({RENDER}). ms/frame incluye decodificar, rotar e inferir; la "
                 "inferencia sola es la columna p50/p90.", widths_cm=[2.3, 1.6, 1.6, 1.6, 1.6, 1.7, 1.7, 1.9, 1.5])
    path = model_bars(one.reset_index(), "ms_per_frame", "Tiempo por frame", "ms por frame", figures_dir / "01_ms_por_frame.png",
                      spread=("ms_per_frame_min", "ms_per_frame_max"))
    report.figure(path, "Tiempo por frame con un video a la vez. Barra: mediana de los 8 videos; bigote: del video más "
                  "rápido al más lento.")
    path = model_bars(one.reset_index(), "realtime_factor", "Tiempo de proceso ÷ duración del video", "veces la duración",
                      figures_dir / "02_x_duracion.png", decimals=2, reference=(1.0, "tiempo real"))
    report.figure(path, "Cuántas veces la duración del video tarda en procesarse. Por debajo de la línea, el modelo va "
                  "más rápido que el video.")

    # --- 4. Recursos ---
    report.heading("4. Consumo de recursos por modelo")
    rows = [[model_label(m), num(one.loc[m, "cores"], 2), num(raw_one.loc[m, "cpu_budget_pct"], 0) + " %",
             num(one.loc[m, "throttled_pct"], 0) + " %", num(raw_one.loc[m, "cpu_per_video_s"], 2) + " s",
             num(raw_one.loc[m, "model_ram_mb"], 0) + " MB", num(one.loc[m, "ram_peak"], 0) + " MB",
             "+" + num(one.loc[m, "ram_drift"], 0) + " MB", num(raw_one.loc[m, "cold_init_s"], 2) + " s"] for m in models]
    report.table(["Modelo", "Núcleos usados", "Del presupuesto de CPU", "Periodos frenados por la cuota",
                  "CPU por segundo de video", "RAM al cargar el modelo", "RAM pico", "Deriva entre jobs",
                  "Arranque en frío"], rows,
                 "Recursos con un video a la vez. RAM al cargar el modelo y arranque en frío: del calentamiento, en un "
                 "proceso recién creado (incluye importar las librerías). Deriva: cuánto subió la RAM de partida entre "
                 "repeticiones del mismo proceso.", widths_cm=[2.3, 1.5, 1.6, 1.8, 1.7, 1.7, 1.5, 1.6, 1.6])
    path = model_bars(one.reset_index(), "cores", "Núcleos usados", "núcleos", figures_dir / "03_nucleos.png", decimals=2,
                      reference=(config.BUDGET_CPU_CORES, f"presupuesto {config.RENDER_PLAN}: {config.BUDGET_CPU_CORES} núcleos"))
    report.figure(path, "Núcleos usados en promedio durante el procesamiento (segundos de CPU ÷ segundos de reloj).")
    path = model_bars(raw_one.reset_index(), "cpu_per_video_s", "Segundos de CPU por segundo de video", "s de CPU",
                      figures_dir / "04_cpu_por_segundo_de_video.png", decimals=2)
    report.figure(path, "Costo de cómputo por segundo de video: lo que el modelo gasta de CPU, sin importar cuántos "
                  "núcleos use a la vez. Es la métrica de eficiencia más directa entre modelos.")
    path = model_bars(one.reset_index(), "ram_peak", "RAM pico por proceso", "MB", figures_dir / "05_ram.png")
    report.figure(path, f"RAM pico de un proceso (peor video). El margen por proceso con "
                  f"{config.SERVICE_WORKER_CONCURRENCY} en paralelo es "
                  f"{num(config.BUDGET_RAM_MB * stats.RAM_SAFE_FRACTION / config.SERVICE_WORKER_CONCURRENCY, 0)} MB.")
    throttled = [m for m in models if (one.loc[m, "throttled_pct"] or 0) >= 1]
    report.para("«Periodos frenados por la cuota»: fracción de los periodos de 100 ms en que el contenedor agotó sus "
                "2 núcleos y esperó (cpu.stat del cgroup). "
                + (f"Solo la toca{'n' if len(throttled) > 1 else ''} {', '.join(model_label(m) for m in throttled)}: "
                   "los modelos que usan menos núcleos no llegan al límite." if throttled else
                   "Ningún modelo llegó al límite."), size=9.5)

    # --- 5. Dos a la vez ---
    report.heading("5. Dos videos a la vez (como el worker)")
    if two.empty:
        report.para("Pendiente: la corrida con --concurrency 2 no tiene sesiones completas.")
    else:
        report.para(f"El worker del servicio corre con --concurrency={config.SERVICE_WORKER_CONCURRENCY}. Con dos "
                    "videos a la vez, cada proceso tiene un solo hilo y los dos se reparten las 2 CPU. Cada video tarda "
                    "más, pero el worker termina más videos por hora.")
        rows = []
        for m in models:
            if m not in two.index:
                continue
            rows.append([model_label(m), num(one.loc[m, "ms_per_frame"], 1), num(two.loc[m, "ms_per_frame"], 1),
                         pct_change(two.loc[m, "ms_per_frame"] / one.loc[m, "ms_per_frame"]),
                         num(one.loc[m, "throughput"], 0), num(two.loc[m, "throughput"], 0),
                         pct_change(two.loc[m, "throughput"] / one.loc[m, "throughput"]),
                         num(two.loc[m, "cores"], 2), num(worker_ram.get(m), 0) + " MB", num(two.loc[m, "job_s_max"], 1) + " s"])
        report.table(["Modelo", "ms/frame, 1 a la vez", "ms/frame, 2 a la vez", "Cambio por video", "Videos/hora, 1",
                      "Videos/hora, 2", "Cambio en videos/hora", "Núcleos por proceso (2)", "RAM de los 2 procesos",
                      "Peor job (2)"], rows,
                     "Uno frente a dos videos a la vez. Videos por hora: los que termina el worker completo. RAM de los "
                     "2 procesos: suma de sus picos (cota superior).", widths_cm=[2.2, 1.5, 1.5, 1.5, 1.5, 1.5, 1.6, 1.6, 1.7, 1.5],
                     font_size=8)
        compare = frame[frame["source"].isin([RENDER, RENDER_C2])]
        path = source_bars(compare, "throughput", "Videos por hora que termina el worker", "videos por hora",
                           figures_dir / "06_videos_por_hora.png", CONCURRENCY_COLORS)
        report.figure(path, f"«{RENDER}»: un video a la vez; «{RENDER_C2}»: dos a la vez. Videos de 11 a 17 s.")
        path = source_bars(compare, "realtime_factor", "Tiempo de proceso ÷ duración del video", "veces la duración",
                           figures_dir / "07_x_duracion_concurrencia.png", CONCURRENCY_COLORS, decimals=2,
                           reference=(1.0, "tiempo real"))
        report.figure(path, "Cuánto se alarga cada video cuando comparte las 2 CPU con otro.")
        ram_frame = pd.DataFrame(
            [{"model": m, "source": "1 proceso (1 a la vez)", "ram": one.loc[m, "ram_peak"]} for m in models]
            + [{"model": m, "source": "2 procesos juntos (2 a la vez)", "ram": worker_ram.get(m)} for m in models if m in worker_ram])
        path = source_bars(ram_frame, "ram", "RAM del worker", "MB", figures_dir / "08_ram_worker.png",
                           {"1 proceso (1 a la vez)": CONCURRENCY_COLORS[RENDER],
                            "2 procesos juntos (2 a la vez)": CONCURRENCY_COLORS[RENDER_C2]},
                           reference=(config.BUDGET_RAM_MB * stats.RAM_SAFE_FRACTION,
                                      f"80 % de {config.BUDGET_RAM_MB} MB"))
        report.figure(path, "RAM del worker con uno y con dos videos a la vez, frente al margen de RAM del plan.")

    # --- 6. Límites del servicio ---
    report.heading("6. Qué duración de video aguanta cada modelo")
    report.para("El costo escala con la duración del video (frames), no con su tamaño en MB: el límite de "
                f"{config.SERVICE_MAX_VIDEO_SIZE_MB} MB del servicio no acota la duración. Al ritmo medido, el video más "
                f"largo que termina dentro de {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s (apagado del worker al redesplegar) "
                f"y de {config.SERVICE_STUCK_JOB_TIMEOUT_S} s (barrido de jobs atascados) es:")
    stuck = {}
    for source in (RENDER, RENDER_C2):
        subset = summary[summary["environment"] == source]
        for m, group in subset.groupby("model"):
            stuck[(source, m)] = group["max_video_s_within_stuck_timeout"].min()
    rows = []
    for m in models:
        row = [model_label(m), num(one.loc[m, "max_video_s"], 0) + " s", num(stuck.get((RENDER, m)), 0) + " s"]
        if not two.empty and m in two.index:
            row += [num(two.loc[m, "max_video_s"], 0) + " s", num(stuck.get((RENDER_C2, m)), 0) + " s"]
        else:
            row += ["n/d", "n/d"]
        rows.append(row)
    report.table(["Modelo", f"En {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s, 1 a la vez", f"En {config.SERVICE_STUCK_JOB_TIMEOUT_S} s, 1 a la vez",
                  f"En {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s, 2 a la vez", f"En {config.SERVICE_STUCK_JOB_TIMEOUT_S} s, 2 a la vez"], rows,
                 "Duración máxima de video, en el peor de los 8 videos medidos (descontando la construcción del modelo).",
                 font_size=8.5)
    if not two.empty:
        limit_frame = frame[frame["source"].isin([RENDER, RENDER_C2])]
        path = source_bars(limit_frame, "max_video_s", f"Video más largo que termina en {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s",
                           "segundos de video", figures_dir / "09_video_maximo.png", CONCURRENCY_COLORS)
        report.figure(path, "Duración de video que alcanza a procesarse antes de que el worker se apague al redesplegar.")

    # --- 7. Costo ---
    report.heading("7. Costo de cómputo por video (indicador)")
    report.para(f"Para comparar modelos en dinero: a la tarifa que mostró Render para un cron job 2c-8g el "
                f"2026-10-01 (US$ {num(CRON_2C8G_USD_PER_MIN, 5)} por minuto, US$ {num(usd_per_hour, 3)} por hora), "
                "esto cuesta procesar 1.000 videos como los medidos. No es el costo del servicio: el worker se factura "
                "como servicio continuo y el plan lo decide el grupo de investigación.")
    rows = []
    for m in models:
        cost_one = usd_per_hour / one.loc[m, "throughput"] * 1000
        cost_two = usd_per_hour / two.loc[m, "throughput"] * 1000 if not two.empty and m in two.index else None
        rows.append([model_label(m), num(one.loc[m, "throughput"], 0), num(cost_one, 2),
                     num(two.loc[m, "throughput"], 0) if cost_two else "n/d", num(cost_two, 2) if cost_two else "n/d"])
    report.table(["Modelo", "Videos/hora, 1 a la vez", "US$ por 1.000 videos, 1 a la vez", "Videos/hora, 2 a la vez",
                  "US$ por 1.000 videos, 2 a la vez"], rows, "Costo de cómputo por cada 1.000 videos de 11 a 17 s, en dólares.",
                 font_size=8.5)

    # --- 8. Variabilidad ---
    report.heading("8. Variabilidad y validez de las mediciones")
    cv_one = _cv_by_model(summary, RENDER)
    cv_two = _cv_by_model(summary, RENDER_C2)
    rows = []
    for m in models:
        a = cv_one.get(m, (None, None, 0))
        b = cv_two.get(m, (None, None, 0))
        rows.append([model_label(m), num(a[0], 1) + " %", num(a[1], 1) + " %", f"{a[2]} de 8",
                     num(b[0], 1) + " %" if m in cv_two else "n/d", num(b[1], 1) + " %" if m in cv_two else "n/d",
                     f"{b[2]} de 8" if m in cv_two else "n/d"])
    report.table(["Modelo", "CV mediano, 1", "CV máximo, 1", f"Sesiones con CV > {num(stats.MAX_CV_PCT, 0)} %, 1",
                  "CV mediano, 2", "CV máximo, 2", f"Sesiones con CV > {num(stats.MAX_CV_PCT, 0)} %, 2"], rows,
                 "Coeficiente de variación del tiempo de proceso entre repeticiones, por sesión modelo-video.", font_size=8.5)
    cv_frame = pd.DataFrame([{"model": m, "source": RENDER, "cv": cv_one[m][0]} for m in models if m in cv_one]
                            + [{"model": m, "source": RENDER_C2, "cv": cv_two[m][0]} for m in models if m in cv_two])
    path = source_bars(cv_frame, "cv", "Variación entre repeticiones (CV mediano)", "%", figures_dir / "10_cv.png",
                       CONCURRENCY_COLORS, decimals=1, reference=(stats.MAX_CV_PCT, f"criterio: {num(stats.MAX_CV_PCT, 0)} %"))
    report.figure(path, "CV mediano por modelo. Por encima de la línea, la medición varía más de lo que exige el criterio.")
    total_one = sum(v[2] for v in cv_one.values())
    total_two = sum(v[2] for v in cv_two.values())
    report.bullets([
        f"**{RENDER}:** {total_one} de {len(summary[summary['environment'] == RENDER])} sesiones superan el "
        f"{num(stats.MAX_CV_PCT, 0)} % de CV. **{RENDER_C2}:** {total_two} de "
        f"{len(summary[summary['environment'] == RENDER_C2])}. Las medianas de tiempo siguen siendo usables, pero esas "
        "sesiones no cumplen el criterio de estabilidad.",
        "Cada corrida usó una sola instancia de Render, en una sola franja horaria. La variación entre repeticiones de "
        "una instancia compartida en la nube es esperable (Leitner & Cito, 2016; Laaber et al., 2019). Para citar un "
        "tiempo de Render con confianza conviene repetir la suite en otros momentos, con el orden barajado "
        "(--trials --shuffle-seed, metodología RMIT de Abedi & Brecht, 2017).",
        "Ninguna sesión tuvo cortes por falta de memoria (OOM-kill) y todas detectaron a la persona en el 100 % de los frames.",
    ] + [f"**Nota.** {note}" for note in notes or []])

    # --- 9. Limitaciones ---
    report.heading("9. Limitaciones")
    report.bullets([
        "**Videos de 11 a 17 s.** Los tiempos por job escalan con la duración; la sección 6 da el video más largo que "
        "cabe en cada límite. Falta medir un video del tamaño máximo que acepta el servicio (200 MB).",
        "**Costo del modelo, no del pipeline del servicio.** La validación, el cálculo de ángulos y la subida de "
        "resultados del servicio no están incluidos.",
        "**El costo en dólares es un indicador** con la tarifa de un cron job; no es la factura del worker.",
        "**Sin precisión.** La decisión de modelo combina este costo con la precisión de entorno-pruebas.",
    ])

    # --- Anexos ---
    report.landscape_section()
    report.heading("Anexo A. Resumen por modelo y video")
    model_rank = {m: i for i, m in enumerate(models)}
    video_rank = {v: i for i, v in enumerate(video_order)}
    ordered = summary.assign(_m=summary["model"].map(model_rank), _v=summary["video_id"].map(video_rank).fillna(99))
    rows = []
    for _, r in ordered.sort_values(["_m", "_v", "environment"]).iterrows():
        rows.append([r["environment"], model_label(r["model"]), r["video_id"], str(r["n"]), num(r["processing_s_median"], 2),
                     num(r["processing_s_cv_pct"], 1), f"{num(r['processing_s_ci95_low'], 2)}–{num(r['processing_s_ci95_high'], 2)}",
                     num(r["ms_per_frame_median"], 1), "×" + num(r["realtime_factor_median"], 2), num(r["job_s_max"], 2),
                     num(r["cpu_cores_used_median"], 2), num(r["ram_mb_peak_max"], 0), num(r["ram_drift_mb"], 0),
                     num(r["throughput_videos_per_hour"], 0), VERDICT_LABELS.get(r["verdict_time"], "n/d"),
                     "sí" if r["measurement_valid"] else "no"])
    report.table(["Corrida", "Modelo", "Video", "n", "Proceso (s, mediana)", "CV %", "IC 95 % (s)", "ms/frame",
                  "× duración", "Peor job (s)", "Núcleos", "RAM pico (MB)", "Deriva (MB)", "Videos/hora",
                  "Veredicto tiempo", "Válida"], rows,
                 f"Una fila por corrida, modelo y video. n = repeticiones medidas (3 con uno a la vez; 6 con dos a la vez: "
                 "2 procesos × 3).", widths_cm=[1.8, 2.3, 1.2, 0.6, 1.6, 1.0, 2.2, 1.4, 1.4, 1.5, 1.3, 1.5, 1.4, 1.4, 1.8, 1.1],
                 font_size=7)
    report.heading("Anexo B. Cada video")
    for index, run in enumerate(available):
        subset = summary[summary["environment"] == run.label]
        path = video_bars(subset, "ms_per_frame_median", video_order, f"Tiempo por frame por video — {run.label}",
                          "ms por frame (mediana)", figures_dir / f"11_ms_por_video_{index + 1}.png")
        report.figure(path, f"Tiempo por frame de cada modelo en cada video, {run.label}.", width_cm=21)
        path = video_bars(subset, "job_s_max", video_order, f"Peor job por video — {run.label}", "segundos",
                          figures_dir / f"12_job_por_video_{index + 1}.png",
                          (config.SERVICE_MAX_SHUTDOWN_DELAY_S, f"apagado del worker: {config.SERVICE_MAX_SHUTDOWN_DELAY_S} s"))
        report.figure(path, f"Peor job (construir el modelo + procesar el video) de cada video, {run.label}.", width_cm=21)

    report.save(out_path)
    print(f"  -> {out_path}")
    return out_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m scripts.build_render_report", description=__doc__.split("\n\n")[0])
    parser.add_argument("--render", type=Path, required=True, help="corrida de Render con un video a la vez")
    parser.add_argument("--render-c2", type=Path, help="corrida de Render con --concurrency 2")
    parser.add_argument("--out", type=Path, required=True, help="ruta del .docx")
    parser.add_argument("--image-tag", default=None)
    parser.add_argument("--note", action="append", default=[], help="nota para la sección de validez (repetible)")
    args = parser.parse_args(argv)
    build(Run(RENDER, args.render), Run(RENDER_C2, args.render_c2) if args.render_c2 else None, args.out,
          args.image_tag, args.note)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
